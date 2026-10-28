"""Bygger utskicksfilen för lead-mejl sats 2 (Kent, Niclas och Johan).

Läser Leadslista_kompletterad.xlsx, behåller flikarna Leadsbas och Kontroll
orörda och lägger till: Läs mig, en flik per avsändare, Uppvärmning
och Antaganden. Utskicksflikarna har rubriken på rad 1 så att de går att
importera direkt i /portal/lead-mejl.

    python3 docs/leads/build_utskick.py IN.xlsx UT.xlsx \
        [--berikning berikning.jsonl ...] [--kontakter Niclas=fil.xlsx] [--newwave per_rad.json]

--berikning läser utdata från enrich_members.py (webbplats, e-post och
medlemsantal med källa). --kontakter läser en säljares egen kontaktlista
(Kundnamn, Kontaktperson, Telefonnummer, E-post).
"""

from __future__ import annotations

import argparse
import json
import re
import subprocess
import sys
from collections import defaultdict
from concurrent.futures import ThreadPoolExecutor
from datetime import date
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

# Föreningar utan "Säljare Roots" får introduktionen från Johan (HK).
SELLERS = ("Kent", "Niclas", "Johan")
CENTRAL = "Johan"
PACKS, PACK_PRICE, CLUB_SHARE = 6, 399, 0.35
BIG_CLUB = 1000

# Leadsbas, kolumnindex (0-baserat), version 2 av listan (21 kolumner).
C_NAME, C_LAN, C_KOMMUN, C_ORG, C_SPORT, C_SELLER = 0, 1, 2, 3, 4, 5
C_SURPLUS, C_KUVERT, C_ORGNR, C_SITE, C_MEMBERS, C_YEAR = 6, 7, 8, 9, 10, 11
C_CONTACT, C_ROLE = 12, 13
C_EMAILS = (14, 15, 16, 17)
C_PHONE, C_SKIP, C_CATEGORY = 18, 19, 20
EXPECTED_HEADERS = {C_SELLER: "Säljare Roots", C_EMAILS[0]: "E-post 1", C_CATEGORY: "Kategori"}

EMAIL_RE = re.compile(r"^[^@\s,;<>]+@[^@\s,;<>]+\.[a-z]{2,}$", re.I)

BROWN = "5B3A29"
INK = "2B2118"
SAND = "F4EEE6"
FILL_ME = "FFF2C2"
TRACK = "E6EEE8"
READY = "DDEFD9"

SEG_PERSONAL = "Personligt mejl"
SEG_GENERAL = "Generellt mejl"
SEG_DIALOG = "Pågående dialog – påstötning"
SEG_AGREEMENT = "Avtal – skicka inte"
SEG_SKIP = "Skicka inte"
SEG_TEAM = "Lag – eget utskick"
SEGMENTS = [SEG_PERSONAL, SEG_GENERAL, SEG_DIALOG, SEG_TEAM, SEG_AGREEMENT, SEG_SKIP]
SEGMENT_FILL = {
    SEG_PERSONAL: "DDEFD9",
    SEG_GENERAL: "EEE6DA",
    SEG_DIALOG: "F6D3A8",
    SEG_AGREEMENT: "D9D9D9",
    SEG_SKIP: "D9D9D9",
    SEG_TEAM: "D6E4F0",
}
NAME_FIXED = "FBE3D6"
TEAM_SHEET = "Lag"

STATUSES = [
    "Ej skickad", "Skickad", "Öppnad", "Öppnat förslaget", "Vill veta mer",
    "Svarat", "Möte bokat", "Inte intresserad", "Studsade", "Avregistrerad",
]


def text(v) -> str:
    if v is None:
        return ""
    s = str(v).strip()
    return "" if s in ("0", "None") else s


# Säljarnas anteckningar i namnet. Slutat/upphört tas bort ur utskicket, övriga flaggas.
STOPPED = re.compile(r"\bslutat?\b|upphör", re.I)
DORMANT = re.compile(r"\b(konkurs|vilande|inaktiv|kundförl\w*)\b", re.I)
STATUS_SUFFIX = re.compile(r"\s*/\s*(konkurs|vilande|inaktiv|kundförl\w*)\s*$", re.I)
PARENS = re.compile(r"\s*\([^)]*\)?")
# Ordet lag (Lag: P-07, laget, damlaget, A-lag) och lagbeteckningar som (Team P05), (P-05), (F-06/07).
TEAM_WORD = re.compile(r"(?<![\wåäö])(\w*-)?(dam|herr|flick|pojk|ungdoms|junior)?lag(et|en)?(?![\wåäö])", re.I)
TEAM_PARENS = re.compile(r"\(\s*(team|trupp|grupp)\b|\(\s*[PFU][\s/-]*\d|\(\s*(dam|flickor|pojkar)\s*\)", re.I)


def clean_name(name: str) -> str:
    """Namnet utan parenteser och statusanteckningar. Originalet i Förening lämnas orört."""
    s = PARENS.sub("", name)
    s = STATUS_SUFFIX.sub("", s)
    s = re.sub(r"\s+", " ", s).strip(" -/,.")
    return s or name.strip()


def team_reason(name: str) -> str:
    if TEAM_WORD.search(name):
        return "Ordet lag i namnet"
    if TEAM_PARENS.search(name):
        return "Lagbeteckning i namnet (Team, P-05, F-06)"
    return ""


def emails_of(row) -> list[str]:
    out: list[str] = []
    for i in C_EMAILS:
        for part in re.split(r"[;,\s/]+", text(row[i])):
            p = part.lower().removeprefix("mailto:")
            if EMAIL_RE.match(p) and p not in out:
                out.append(p)
    return out


def site_of(row) -> str:
    s = text(row[C_SITE])
    return s if re.match(r"^(https?://|www\.)", s, re.I) else ""


def members_of(row) -> int | None:
    v = row[C_MEMBERS]
    return v if isinstance(v, int) and v > 0 else None


def number(v) -> float:
    return float(v) if isinstance(v, (int, float)) else 0.0


def collect(ws):
    header = [c.value for c in ws[1]]
    for i, name in EXPECTED_HEADERS.items():
        if header[i] != name:
            raise SystemExit(f"Kolumn {i + 1} heter {header[i]!r}, väntade {name!r}. Fel version av listan?")
    leads = {s: [] for s in SELLERS}
    skipped = []
    stopped = []
    for rownum, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not text(row[C_NAME]):
            continue
        seller = text(row[C_SELLER]) or CENTRAL
        if seller not in leads:
            raise SystemExit(f"Okänd säljare {seller!r} på rad {rownum}.")
        if text(row[C_SKIP]):
            skipped.append((seller, text(row[C_NAME]), rownum))
            continue
        mails = emails_of(row)
        if seller == CENTRAL and not mails:
            continue
        if STOPPED.search(text(row[C_NAME])):
            stopped.append((seller, text(row[C_NAME]), rownum))
            continue
        dormant = DORMANT.search(text(row[C_NAME]))
        leads[seller].append({
            "row": rownum,
            "seller": seller,
            "name": text(row[C_NAME]),
            "clean": clean_name(text(row[C_NAME])),
            "fixed": "",
            "fixed_why": "",
            "team": team_reason(text(row[C_NAME])),
            "contact": text(row[C_CONTACT]),
            "email": mails[0] if mails else "",
            "more": ", ".join(mails[1:]),
            "site": site_of(row),
            "phone": text(row[C_PHONE]),
            "kommun": text(row[C_KOMMUN]),
            "lan": text(row[C_LAN]).title(),
            "sport": text(row[C_SPORT]),
            "orgnr": text(row[C_ORGNR]),
            "members": members_of(row),
            "year": row[C_YEAR] if isinstance(row[C_YEAR], int) and row[C_YEAR] > 0 else None,
            "surplus": number(row[C_SURPLUS]),
            "kuvert": number(row[C_KUVERT]),
            "category": text(row[C_CATEGORY]).upper(),
            "check": (f"Namnet säger '{dormant.group(1)}'. Kontrollera att föreningen finns kvar "
                      "innan den får mejl, töm sedan rutan.") if dormant else "",
            "contact_from": "Listan" if text(row[C_CONTACT]) else "",
            "total": members_of(row),
            "active": None,
            "member_status": "Från listan" if members_of(row) else "",
            "member_source": "",
            "member_quote": "",
            "socials": "",
            "shop": "",
            "founded": "",
            "about": "",
        })
    return leads, skipped, stopped


NOT_SHOP = re.compile(r"/butiker/|gclid|gad_source|(^|\.)(ica|coop|willys|hemkop|lidl|happyhomes|bolist)\.se", re.I)
PLATFORM_MAIL = re.compile(r"@(laget|svenskalag|idrottonline|klubbenonline|sportadmin|ibis)\.", re.I)
NEWWAVE_MIN_CONFIDENCE = 60


def add_extras(lead, socials=None, shop=None, founded=None, about=None) -> None:
    if socials and not lead["socials"]:
        lead["socials"] = "\n".join(socials)
    if shop and not lead["shop"] and not NOT_SHOP.search(shop):
        lead["shop"] = shop
    if founded and not lead["founded"] and re.fullmatch(r"(18|19|20)\d\d", str(founded)) \
            and int(founded) <= date.today().year:
        lead["founded"] = str(founded)
    if about and not lead["about"] and len(about) > 25:
        lead["about"] = about[:300]


def apply_newwave(leads, path: str) -> int:
    """Newwaves klubbregister (gratis skrapning): webbplats, e-post, sociala medier, webbshop."""
    data = {int(k): v for k, v in json.load(open(path, encoding="utf-8")).items()}
    seen = defaultdict(int)
    for v in data.values():
        for m in v.get("emails") or []:
            seen[m.lower()] += 1
    n = 0
    for seller in SELLERS:
        for lead in leads[seller]:
            v = data.get(lead["row"])
            if not v or (v.get("confidence") or 0) < NEWWAVE_MIN_CONFIDENCE:
                continue
            n += 1
            if v.get("url") and not lead["site"]:
                lead["site"] = v["url"]
            mails = [m.lower() for m in v.get("emails") or []
                     if EMAIL_RE.match(m) and seen[m.lower()] <= 2 and not PLATFORM_MAIL.search(m)]
            known = [m for m in [lead["email"], *lead["more"].split(", ")] if m]
            for m in mails:
                if m in known:
                    continue
                if not lead["email"] and not lead["check"]:
                    lead["email"] = m
                else:
                    lead["more"] = ", ".join(x for x in [lead["more"], m] if x)
                known.append(m)
            socials = v.get("socials") or {}
            add_extras(lead, socials=list(socials.values()) if isinstance(socials, dict) else socials,
                       shop=v.get("shop"), founded=v.get("founded"), about=v.get("description"))
    return n


def norm_name(s: str) -> str:
    return re.sub(r"[^a-zåäö0-9]", "", s.lower())


def apply_contacts(leads, seller: str, path: str) -> int:
    """Säljarens egen kontaktlista: kontaktperson, telefon och hans e-post först."""
    rows = load_workbook(path, read_only=True).worksheets[0].iter_rows(min_row=2, values_only=True)
    mine = leads[seller]
    by_name = defaultdict(list)
    by_mail = defaultdict(list)
    for lead in mine:
        by_name[norm_name(lead["name"])].append(lead)
        for m in [lead["email"], *lead["more"].split(", ")]:
            if m:
                by_mail[m].append(lead)
    n = 0
    for r in rows:
        name, person, phone, mails = text(r[3]), text(r[4]), text(r[5]), text(r[6])
        if not name:
            continue
        own = [m.strip().lower() for m in re.split(r"[;,\s]+", mails) if EMAIL_RE.match(m.strip())]
        hits = by_name.get(norm_name(name)) or next((by_mail[m] for m in own if m in by_mail), [])
        if not hits:
            continue
        lead = sorted(hits, key=lambda l: l["category"] != "A")[0]
        n += 1
        lead["category"] = "A"
        if person:
            names = []
            for p in person.split(";"):
                last, _, first = p.strip().partition(", ")
                names.append(f"{first} {last}".strip() if first else p.strip())
            lead["contact"] = "; ".join(n for n in names if n)
            owner = seller if seller.endswith("s") else f"{seller}s"
            lead["contact_from"] = f"{owner} lista 2021 – bekräfta att personen är kvar"
        if phone and not lead["phone"]:
            lead["phone"] = phone
        if own:
            known = [m for m in [lead["email"], *lead["more"].split(", ")] if m]
            ordered = own + [m for m in known if m not in own]
            lead["email"], lead["more"] = ordered[0], ", ".join(ordered[1:])
    return n


def normalize_record(rec: dict) -> dict:
    """enrich_cheap.py sparar egen sidläsning, modell och sökning var för sig; gör om till enrich_members-form."""
    if "found" in rec:
        return rec
    free, model = rec.get("free") or {}, rec.get("model") or {}
    search = rec.get("search") or {}
    if search.get("found"):
        found, check = dict(search["found"]), search.get("check") or {}
    else:
        found = {k: model.get(k) for k in (
            "members_total", "members_active", "members_year", "members_quote", "members_source")}
        found.update(website=free.get("site"), email=None, note=None)
        check = {}
    if not found.get("website"):
        found["website"] = free.get("site")
    if not found.get("email") and free.get("emails"):
        found["email"], check["email_mx"], check["email_on_page"] = free["emails"][0], True, True
    contact = model.get("contact_name")
    if contact and model.get("contact_role"):
        contact = f"{contact} ({model['contact_role']})"
    return {"row": rec["row"], "found": found, "check": check, "extras": {
        "socials": free.get("socials"), "shop": (free.get("shops") or [None])[0],
        "founded": free.get("founded"), "about": model.get("description"), "contact": contact}}


def apply_enrichment(leads, path: str) -> int:
    """Agentens sökning: webbplats, e-post och medlemsantal med källa och status."""
    by_row = {l["row"]: l for s in SELLERS for l in leads[s]}
    n = 0
    for line in open(path, encoding="utf-8"):
        if not line.strip():
            continue
        rec = normalize_record(json.loads(line))
        lead = by_row.get(rec["row"])
        if not lead:
            continue
        n += 1
        f, ck = rec["found"], rec["check"]
        x = rec.get("extras") or {}
        add_extras(lead, socials=x.get("socials"), shop=x.get("shop"),
                   founded=x.get("founded"), about=x.get("about"))
        if x.get("contact") and not lead["contact"]:
            lead["contact"] = x["contact"]
            lead["contact_from"] = "Agenten, föreningens webbplats"
        if f.get("website") and not lead["site"]:
            lead["site"] = f["website"]
        mail = (f.get("email") or "").strip().lower()
        if mail and ck.get("email_mx") and EMAIL_RE.match(mail):
            known = [m for m in [lead["email"], *lead["more"].split(", ")] if m]
            if not lead["email"] and ck.get("email_on_page"):
                lead["email"] = mail
            elif mail not in known:
                lead["more"] = ", ".join(x for x in [lead["more"], mail] if x)
        if lead["member_status"]:
            continue
        total, active, year = f.get("members_total"), f.get("members_active"), f.get("members_year")
        if not (total or active):
            continue
        lead["total"], lead["active"] = total, active
        lead["year"] = year
        lead["member_source"] = f.get("members_source") or ""
        lead["member_quote"] = " ".join(x for x in [f.get("members_quote"), f.get("note")] if x)
        value = active or total
        if (active and total and active > total) or value < 50:
            lead["member_status"] = "Misstänkt fel"
            lead["check"] = " ".join(x for x in [
                f"Agenten hittade {total or '–'} medlemmar ({active or '–'} aktiva), som ser fel ut. "
                "Fyll i rätt antal deltagare och töm rutan.", lead["check"]] if x)
            continue
        lead["members"] = value
        if year and year >= 2023 and ck.get("members_on_page"):
            lead["member_status"] = "Agenten – källa kontrollerad"
        elif year and year >= 2023:
            lead["member_status"] = "Agenten – kontrollera källan"
        else:
            lead["member_status"] = "Agenten – gammal siffra"
            lead["check"] = " ".join(x for x in [
                f"Medlemsantalet är från {year or 'okänt år'}. Bekräfta antal deltagare och töm rutan.",
                lead["check"]] if x)
    return n


def mark_shared(leads) -> int:
    """Samma adress för flera föreningar: behåll den på en rad, spärra resten."""
    by_mail = defaultdict(list)
    for seller in SELLERS:
        for lead in leads[seller]:
            if lead["email"]:
                by_mail[lead["email"]].append(lead)
    shared = 0
    for mail, group in by_mail.items():
        if len(group) < 2:
            continue
        shared += 1
        group.sort(key=lambda l: (l["category"] != "A", l["members"] is None, -l["surplus"]))
        keeper = group[0]
        for other in group[1:]:
            other["check"] = " ".join(x for x in [
                f"Samma adress ({mail}) som {keeper['name']}, rad {keeper['row']} i Leadsbas. "
                "Hitta egen adress eller töm rutan om ett gemensamt mejl räcker.",
                other["check"],
            ] if x)
            other["email"] = ""
            other["more"] = ", ".join(x for x in [mail, other["more"]] if x)
    return shared


def mark_big(leads) -> None:
    for seller in SELLERS:
        for lead in leads[seller]:
            m = lead["members"]
            if m is not None and m > BIG_CLUB and not lead["check"]:
                lead["check"] = (
                    f"Stor förening ({m:,} medlemmar). Bedöm antal deltagare "
                    "och töm sedan den här rutan.".replace(",", " ")
                )


def has_mx(domain: str) -> bool:
    out = subprocess.run(
        ["dig", "+short", "+time=3", "+tries=2", "MX", domain],
        capture_output=True, text=True, check=False,
    ).stdout.strip()
    return bool(out)


def mark_dead_domains(leads) -> int:
    domains = {l["email"].split("@")[1] for s in SELLERS for l in leads[s] if l["email"]}
    with ThreadPoolExecutor(max_workers=32) as pool:
        dead = {d for d, ok in zip(domains, pool.map(has_mx, domains)) if not ok}
    n = 0
    for seller in SELLERS:
        for lead in leads[seller]:
            if lead["email"] and lead["email"].split("@")[1] in dead:
                lead["check"] = " ".join(x for x in [
                    f"{lead['email']}: domänen tar inte emot mejl. Hitta en ny adress, "
                    "skriv den i E-post och töm sedan den här rutan.",
                    lead["check"],
                ] if x)
                lead["email"] = ""
                n += 1
    return n


def ready_personal(lead) -> bool:
    return bool(lead["email"]) and lead["members"] is not None and not lead["check"]


def suggested_segment(lead, seller: str) -> str:
    """Förslag; Kent, Niclas och Johan ändrar. Pågående dialog sätts bara av dem."""
    if seller != CENTRAL and (priority(lead) < 3 or lead["category"] == "A"):
        return SEG_PERSONAL
    return SEG_GENERAL


def priority(lead) -> int:
    """1 = klar för personligt mejl och prioriterad (kategori A eller tidigare försäljning)."""
    if ready_personal(lead):
        return 1 if lead["category"] == "A" or lead["surplus"] > 0 else 2
    return 3


HEADERS = [
    # (rubrik, bredd, grupp) grupp: base | fill | calc | track
    ("Förening", 34, "base"),
    ("Korrigerat namn", 30, "fill"),
    ("Namnändring", 30, "base"),
    ("Segment", 26, "fill"),
    ("Kontaktperson", 20, "fill"),
    ("E-post", 32, "fill"),
    ("Fler e-postadresser", 28, "base"),
    ("Webbplats", 30, "fill"),
    ("Antal deltagare", 11, "fill"),
    ("Medlemmar enligt källa", 11, "base"),
    ("Medlemsår", 9, "base"),
    ("Aktiva enligt källa", 10, "base"),
    ("Medlemsuppgift", 24, "base"),
    ("Källa medlemsantal", 30, "base"),
    ("Citat och kommentar", 44, "base"),
    ("Belopp per deltagare", 11, "calc"),
    ("Total försäljning", 13, "calc"),
    ("Behållning 35 %", 13, "calc"),
    ("Klar för intro", 12, "calc"),
    ("Klar för personligt mejl", 22, "calc"),
    ("Att göra", 34, "calc"),
    ("Kontrollera", 40, "fill"),
    ("Ordning", 8, "base"),
    ("Kategori", 8, "base"),
    ("Telefon", 15, "base"),
    ("Kontaktperson från", 22, "base"),
    ("Sociala medier", 30, "base"),
    ("Webbshop", 30, "base"),
    ("Grundår", 8, "base"),
    ("Om föreningen", 44, "base"),
    ("Kommun", 15, "base"),
    ("Län", 14, "base"),
    ("Idrott", 16, "base"),
    ("Organisationsnummer", 14, "base"),
    ("Överskott H20–V21", 12, "base"),
    ("Kuvert H2025", 9, "base"),
    ("Rad i Leadsbas", 9, "base"),
    ("Skickat", 11, "track"),
    ("Öppnat mejlet", 11, "track"),
    ("Öppnat förslaget", 11, "track"),
    ("Vill veta mer", 11, "track"),
    ("Svarat", 11, "track"),
    ("Status", 16, "track"),
    ("Nästa steg", 26, "track"),
    ("Följ upp senast", 12, "track"),
    ("Anteckning", 36, "track"),
]
TEAM_HEADERS = [("Säljare", 10, "base"), ("Varför i Lag-fliken", 30, "base")]
COL = {h: get_column_letter(i + 1) for i, (h, _, _) in enumerate(HEADERS)}

HEAD_FILL = {
    "base": PatternFill("solid", fgColor=BROWN),
    "fill": PatternFill("solid", fgColor="E8B93A"),
    "calc": PatternFill("solid", fgColor="8A6A55"),
    "track": PatternFill("solid", fgColor="4F6B5A"),
}
HEAD_FONT = {
    "base": Font(bold=True, color="FFFFFF"),
    "fill": Font(bold=True, color=INK),
    "calc": Font(bold=True, color="FFFFFF"),
    "track": Font(bold=True, color="FFFFFF"),
}
THIN = Side(style="thin", color="E2D8CC")


def name_change(lead) -> str:
    parts = []
    if lead["clean"] != lead["name"]:
        parts.append("Parentes eller anteckning borttagen")
    if lead["fixed"]:
        parts.append(f"Stavning: {lead['fixed_why']}" if lead["fixed_why"] else "Stavning rättad")
    return ". ".join(parts)


def write_seller_sheet(wb, title: str, seller: str | None, leads: list[dict]):
    """seller=None är Lag-fliken: alla säljares lag samlade, med säljare och skäl längst bak."""
    ws = wb.create_sheet(title)
    headers = HEADERS + (TEAM_HEADERS if seller is None else [])
    for i, (h, width, group) in enumerate(headers, start=1):
        c = ws.cell(1, i, h)
        c.fill, c.font = HEAD_FILL[group], HEAD_FONT[group]
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[1].height = 34

    if seller is None:
        leads = sorted(leads, key=lambda l: (l["clean"].lower(), l["name"].lower()))
    else:
        leads = sorted(leads, key=lambda l: (
            priority(l), l["category"] != "A", not l["email"], -l["surplus"], l["name"].lower()))
    fill_me = PatternFill("solid", fgColor=FILL_ME)
    for r, lead in enumerate(leads, start=2):
        e, n, b, t, k, chk, s = (COL[h] for h in (
            "E-post", "Antal deltagare", "Belopp per deltagare",
            "Total försäljning", "Kontrollera", "Kontaktperson", "Segment"))
        personal = f'{s}{r}="{SEG_PERSONAL}"'
        values = {
            "Förening": lead["name"],
            "Korrigerat namn": lead["fixed"] or lead["clean"],
            "Namnändring": name_change(lead) or None,
            "Segment": SEG_TEAM if seller is None else suggested_segment(lead, seller),
            "Säljare": lead["seller"],
            "Varför i Lag-fliken": lead["team"],
            "Kontaktperson": lead["contact"] or None,
            "E-post": lead["email"] or None,
            "Fler e-postadresser": lead["more"] or None,
            "Webbplats": lead["site"] or None,
            "Antal deltagare": lead["members"],
            "Medlemmar enligt källa": lead["total"],
            "Medlemsår": lead["year"],
            "Aktiva enligt källa": lead["active"],
            "Medlemsuppgift": lead["member_status"] or None,
            "Källa medlemsantal": lead["member_source"] or None,
            "Citat och kommentar": lead["member_quote"] or None,
            "Kontaktperson från": lead["contact_from"] or None,
            "Sociala medier": lead["socials"] or None,
            "Webbshop": lead["shop"] or None,
            "Grundår": int(lead["founded"]) if lead["founded"] else None,
            "Om föreningen": lead["about"] or None,
            "Belopp per deltagare": "=Antaganden!$B$4",
            "Total försäljning": f'=IF({n}{r}="","",{n}{r}*{b}{r})',
            "Behållning 35 %": f'=IF({t}{r}="","",ROUND({t}{r}*Antaganden!$B$5,0))',
            "Klar för intro": f'=IF({e}{r}="","Nej","Ja")',
            "Klar för personligt mejl": (
                f'=IF(NOT({personal}),"–",'
                f'IF({e}{r}="","Nej – saknar e-post",'
                f'IF(OR({n}{r}="",{n}{r}=0),"Nej – saknar antal deltagare",'
                f'IF({k}{r}<>"","Nej – se Kontrollera","Ja"))))'
            ),
            "Att göra": (
                f'=IF(OR({s}{r}="{SEG_AGREEMENT}",{s}{r}="{SEG_SKIP}"),"",'
                f'TRIM(IF({e}{r}="","Hitta e-post. ","")'
                f'&IF(AND({personal},OR({n}{r}="",{n}{r}=0)),"Fyll i antal deltagare. ","")'
                f'&IF({k}{r}<>"","Kolla rutan Kontrollera. ","")'
                f'&IF(AND({personal},{chk}{r}=""),"Valfritt: kontaktperson.","")))'
            ),
            "Kontrollera": lead["check"] or None,
            "Ordning": priority(lead),
            "Kategori": lead["category"] or None,
            "Telefon": lead["phone"] or None,
            "Kommun": lead["kommun"] or None,
            "Län": lead["lan"] or None,
            "Idrott": lead["sport"] or None,
            "Organisationsnummer": lead["orgnr"] or None,
            "Överskott H20–V21": lead["surplus"] or None,
            "Kuvert H2025": lead["kuvert"] or None,
            "Rad i Leadsbas": lead["row"],
            "Status": "Ej skickad",
        }
        for i, (h, _, group) in enumerate(headers, start=1):
            c = ws.cell(r, i, values.get(h))
            c.border = Border(bottom=THIN)
            if group == "fill" and values.get(h) in (None, ""):
                c.fill = fill_me
        if lead["fixed"]:
            ws[f"{COL['Korrigerat namn']}{r}"].fill = PatternFill("solid", fgColor=NAME_FIXED)
        for h in ("Belopp per deltagare", "Total försäljning", "Behållning 35 %", "Överskott H20–V21"):
            ws[f"{COL[h]}{r}"].number_format = '# ##0" kr"'
        for h in ("Skickat", "Följ upp senast"):
            ws[f"{COL[h]}{r}"].number_format = "yyyy-mm-dd"
        ws[f"{COL['Kontrollera']}{r}"].alignment = Alignment(wrap_text=True, vertical="top")
        src = ws[f"{COL['Källa medlemsantal']}{r}"]
        if lead["member_source"].startswith("http"):
            src.hyperlink = lead["member_source"]
            src.font = Font(color="1F5FA8", underline="single")

    last = max(2, len(leads) + 1)
    ws.freeze_panes = f"{COL['Namnändring']}2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{last}"
    seg = DataValidation(type="list", formula1='"' + ",".join(SEGMENTS) + '"', allow_blank=False,
                         showErrorMessage=True, errorTitle="Segment",
                         error="Välj ett segment i listan.")
    ws.add_data_validation(seg)
    seg.add(f"{COL['Segment']}2:{COL['Segment']}{last}")
    for name, color in SEGMENT_FILL.items():
        ws.conditional_formatting.add(
            f"{COL['Segment']}2:{COL['Segment']}{last}",
            FormulaRule(formula=[f'${COL["Segment"]}2="{name}"'], fill=PatternFill("solid", fgColor=color),
                        stopIfTrue=True),
        )
    klar = COL["Klar för personligt mejl"]
    ws.conditional_formatting.add(
        f"A2:{get_column_letter(len(HEADERS))}{last}",
        FormulaRule(formula=[f'${klar}2="Ja"'], fill=PatternFill("solid", fgColor=READY)),
    )
    dv = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{COL['Status']}2:{COL['Status']}{last}")
    yn = DataValidation(type="list", formula1='"Ja,Nej"', allow_blank=True)
    ws.add_data_validation(yn)
    for h in ("Öppnat mejlet", "Öppnat förslaget", "Vill veta mer", "Svarat"):
        yn.add(f"{COL[h]}2:{COL[h]}{last}")
    return leads


def stats(leads):
    def has(l, k):
        return bool(l[k])
    return {
        "total": len(leads),
        "email": sum(has(l, "email") for l in leads),
        "members": sum(l["members"] is not None for l in leads),
        "ready": sum(priority(l) < 3 for l in leads),
        "prio1": sum(priority(l) == 1 for l in leads),
        "cat_a": sum(l["category"] == "A" for l in leads),
        "cat_a_need": sum(l["category"] == "A" and not ready_personal(l) for l in leads),
        "contact": sum(has(l, "contact") for l in leads),
        "site": sum(has(l, "site") for l in leads),
        "no_email": sum(not l["email"] and not l["check"] for l in leads),
        "need_members": sum(bool(l["email"]) and l["members"] is None for l in leads),
        "check": sum(bool(l["check"]) for l in leads),
    }


def write_assumptions(wb):
    ws = wb.create_sheet("Antaganden")
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 70
    ws["A1"] = "Antaganden i räkneexemplet"
    ws["A1"].font = Font(bold=True, size=14, color=BROWN)
    rows = [
        ("Paket per deltagare", PACKS, "Roots Premiumpaket"),
        ("Pris per paket", PACK_PRICE, "kr"),
        ("Belopp per deltagare", "=B2*B3", "[belopp] i mejlet. Ändra här så räknas alla flikar om."),
        ("Andel till föreningen", CLUB_SHARE, "Fast 35 %. Ändras inte utan beslut."),
    ]
    for i, (a, b, c) in enumerate(rows, start=2):
        ws.cell(i, 1, a)
        ws.cell(i, 2, b)
        ws.cell(i, 3, c).font = Font(color="6B5B4E")
    ws["B4"].number_format = '# ##0" kr"'
    ws["B3"].number_format = '# ##0" kr"'
    ws["B5"].number_format = "0 %"
    for cell in ("B2", "B3"):
        ws[cell].fill = PatternFill("solid", fgColor=FILL_ME)
    ws["A7"] = "Mejlets räknerad"
    ws["A7"].font = Font(bold=True)
    ws["A8"] = ("Om [antal deltagare] personer säljer för [belopp] kr var blir försäljningen "
                "[total försäljning] kr. Med 35 % till föreningen ger det [behållning] kr.")
    ws["A8"].alignment = Alignment(wrap_text=True)
    ws.merge_cells("A8:C8")
    ws.row_dimensions[8].height = 36


INTRO_DAYS = [
    (date(2026, 10, 27), 150), (date(2026, 10, 28), 250), (date(2026, 10, 29), 400),
    (date(2026, 10, 30), 600), (date(2026, 11, 2), 800), (date(2026, 11, 3), 1000),
    (date(2026, 11, 4), 1000), (date(2026, 11, 5), 1000),
]
WEEKDAYS = ["mån", "tis", "ons", "tor", "fre", "lör", "sön"]


def intro_schedule(counts: dict[str, int]) -> list[tuple[date, int, dict[str, int]]]:
    """Kents och Niclas föreningar först (i proportion), sedan Johans, inom dagstaket."""
    left = dict(counts)
    days = []
    for day, cap in INTRO_DAYS:
        room, split = cap, {}
        districts = left["Kent"] + left["Niclas"]
        if districts:
            take = min(room, districts)
            kent = min(left["Kent"], round(take * left["Kent"] / districts))
            nic = min(left["Niclas"], take - kent)
            kent = take - nic
            split.update(Kent=kent, Niclas=nic)
            left["Kent"] -= kent
            left["Niclas"] -= nic
            room -= take
        if room and left[CENTRAL]:
            take = min(room, left[CENTRAL])
            split[CENTRAL] = take
            left[CENTRAL] -= take
            room -= take
        days.append((day, cap - room, split))
        if not any(left.values()):
            break
    return days


def write_warmup(wb, intro: dict[str, int], ready: dict[str, int]):
    ws = wb.create_sheet("Uppvärmning")
    widths = {"A": 16, "B": 12, "C": 12, "D": 12, "E": 12, "F": 14, "G": 62}
    for k, v in widths.items():
        ws.column_dimensions[k].width = v
    ws["A1"] = "Uppvärmning av utskicken"
    ws["A1"].font = Font(bold=True, size=14, color=BROWN)
    ws["A2"] = ("Få mejl per dag i början så att roots.nu bygger rykte och hamnar i inkorgen, och så att "
                "Kent, Niclas och Johan hinner ta hand om svaren. Under Resends tak för en ny domän "
                "(150 dag 1, 250, 400, 700, 1 000, 1 500, 2 000).")
    ws["A2"].alignment = Alignment(wrap_text=True)
    ws.merge_cells("A2:G2")
    ws.row_dimensions[2].height = 32

    ws["A4"] = "1. Introduktionsmejlet (samma text till alla, inget räkneexempel)"
    ws["A4"].font = Font(bold=True, color=BROWN)
    head = ["Dag", "Kent", "Niclas", "Johan", "Totalt", "Totalt hittills", "Kommentar"]
    for i, h in enumerate(head, start=1):
        c = ws.cell(5, i, h)
        c.fill, c.font = HEAD_FILL["base"], HEAD_FONT["base"]
    r = 6
    first = r
    for n, (day, total, split) in enumerate(intro_schedule(intro)):
        ws.cell(r, 1, f"{WEEKDAYS[day.weekday()]} {day.day}/{day.month}")
        for col, seller in ((2, "Kent"), (3, "Niclas"), (4, CENTRAL)):
            ws.cell(r, col, split.get(seller) or None)
        ws.cell(r, 5, f"=SUM(B{r}:D{r})")
        ws.cell(r, 6, f"=E{r}" if r == first else f"=F{r - 1}+E{r}")
        note = {0: "Testmejl till egna adresser först. Kolla att det landar i inkorgen."}.get(n, "")
        if not note and split.get(CENTRAL) and not any(split.get(s) for s in ("Kent", "Niclas")):
            note = "Bara Johans föreningar kvar."
        ws.cell(r, 7, note or None)
        r += 1
    ws.cell(r, 1, "Summa").font = Font(bold=True)
    for col in (2, 3, 4, 5):
        L = get_column_letter(col)
        ws.cell(r, col, f"=SUM({L}{first}:{L}{r - 1})").font = Font(bold=True)

    r += 2
    ws.cell(r, 1, "2. Personligt mejl med räkneexempel och shopförslag").font = Font(bold=True, color=BROWN)
    r += 1
    ws.cell(r, 1, ("Startar 5–7 arbetsdagar efter att föreningen fick introduktionen, med prioritet 1 först "
                   "(Niclas kategori A, Kents tidigare försäljning). Bara rader där Klar för personligt mejl = Ja."))
    ws.cell(r, 1).alignment = Alignment(wrap_text=True)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)
    ws.row_dimensions[r].height = 30
    r += 1
    for i, h in enumerate(["Vecka", "Vardagar", "Kent per dag", "Niclas per dag", "Totalt per dag", "Totalt hittills", "Fokus"], start=1):
        c = ws.cell(r, i, h)
        c.fill, c.font = HEAD_FILL["base"], HEAD_FONT["base"]
    plan = [
        ("Från 3/11", 5, 10, 10, "Ring de som inte svarat efter 3–5 arbetsdagar."),
        ("Från 10/11", 5, 20, 20, "Håll studs under 4 %. Pausa om många studsar eller svarar 'sluta'."),
        ("Från 17/11", 5, 30, 30, "Fortsätt med prioritet 2 och de rader som har kompletterats."),
        ("Från 24/11", 5, 40, 40, "40 per dag och säljare är agentens tak idag."),
    ]
    start = r + 1
    for i, (v, days, kent, nic, focus) in enumerate(plan, start=start):
        ws.cell(i, 1, v)
        ws.cell(i, 2, days)
        ws.cell(i, 3, kent).fill = PatternFill("solid", fgColor=FILL_ME)
        ws.cell(i, 4, nic).fill = PatternFill("solid", fgColor=FILL_ME)
        ws.cell(i, 5, f"=C{i}+D{i}")
        ws.cell(i, 6, f"=B{i}*E{i}" if i == start else f"=F{i - 1}+B{i}*E{i}")
        ws.cell(i, 7, focus)
    r = start + len(plan)
    ws.cell(r, 1, f"Klara för personligt mejl nu: Kent {ready['Kent']}, Niclas {ready['Niclas']}.")
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)

    r += 2
    ws.cell(r, 1, "Regler").font = Font(bold=True)
    rules = [
        "Bara vardagar, 08–16.",
        "Studs (felaktig adress) under 4 % och skräppostanmälan under 0,08 % (Resends gränser). Går det över: pausa och rensa listan.",
        "Den som klickar 'Jag vill veta mer' eller svarar kontaktas inom två arbetsdagar.",
        "Den som svarar eller anmäler intresse får inga fler standardmejl. Den som säger nej mejlas aldrig igen.",
        "'Öppnat' är en fingervisning (Apple Mail visar nästan allt som öppnat). Klick och svar är de säkra signalerna.",
    ]
    for i, t in enumerate(rules, start=r + 1):
        ws.cell(i, 1, f"• {t}")
        ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=7)
        ws.cell(i, 1).alignment = Alignment(wrap_text=True)
        ws.row_dimensions[i].height = 30


def apply_spelling(leads, path: str) -> int:
    changes = json.loads(Path(path).read_text())["changes"]
    n = 0
    for lead in (l for s in leads for l in leads[s]):
        fix = changes.get(lead["clean"])
        if fix and fix.get("verified") and fix["fixed"] != lead["clean"]:
            lead["fixed"], lead["fixed_why"] = fix["fixed"], fix["why"]
            n += 1
    return n


def split_teams(leads) -> list[dict]:
    teams = [l for s in leads for l in leads[s] if l["team"]]
    for s in leads:
        leads[s] = [l for l in leads[s] if not l["team"]]
    return teams


def write_readme(wb, st: dict[str, dict], skipped, shared: int, dead: int, stopped, teams, names: dict):
    ws = wb.create_sheet("Läs mig", 0)
    ws.column_dimensions["A"].width = 4
    ws.column_dimensions["B"].width = 110
    k, n, j = st["Kent"], st["Niclas"], st[CENTRAL]
    intro_total = k["email"] + n["email"] + j["email"]
    lines: list[tuple[str, str]] = [
        ("h1", "Leadslista sats 2 – Kent, Niclas och Johan"),
        ("p", f"Byggd {date.today():%Y-%m-%d} från Leadslista_kompletterad.xlsx (version med Säljare Roots och "
              "Kategori). Flikarna Leadsbas och Kontroll är orörda."),
        ("h2", "Utskicket"),
        ("p", "Start måndag 2 november, 10–20 mejl per dag de första två veckorna (uppvärmning av adressen), "
              "sedan ökar vi. Kents och Niclas viktigaste föreningar går först."),
        ("p", "Personligt mejl (räkneexempel och shopförslag) kräver e-post och antal deltagare. Övriga får det "
              "generella mejlet med 'Jag vill veta mer', med Kent eller Niclas som avsändare. Föreningar med avtal "
              "och pågående dialog får inget utskick."),
        ("p", "Fliken Uppvärmning visar den gamla, snabbare takten och gäller inte längre."),
        ("h2", "Segment – vilket utskick föreningen åker med i"),
        ("p", f"Kolumn {COL['Segment']}. Välj i listan; det är den kolumnen systemet läser."),
        ("p", f"{SEG_PERSONAL}: räkneexempel och shopförslag, från Kent eller Niclas. Kräver e-post och antal "
              "deltagare. Förifyllt för Niclas kategori A och föreningar som redan har allt."),
        ("p", f"{SEG_GENERAL}: introduktionen med 'Jag vill veta mer'. Kräver bara e-post. Förifyllt för övriga."),
        ("p", f"{SEG_DIALOG}: föreningar ni redan pratar med. Får inget massutskick, bara påstötningen "
              "('det går bra, dags att bestämma sig') från er egen inkorg. Sätts bara av er."),
        ("p", f"{SEG_TEAM}: förifyllt i fliken Lag. Lagen behandlas för sig och får inte det vanliga utskicket."),
        ("p", f"{SEG_AGREEMENT} och {SEG_SKIP}: får ingenting."),
        ("p", "Ordning (1–3) är något annat: den säger hur komplett underlaget är och styr vem som går först "
              "inom ett segment. 1 = klar för personligt mejl och tidigare försäljning eller kategori A, "
              "2 = klar, 3 = något saknas."),
        ("h2", "Ifyllt av agenten"),
        ("p", "Medlemsuppgift visar var antalet kommer ifrån: 'Från listan', 'Agenten – källa kontrollerad' "
              "(siffran står på källsidan, 2023 eller senare), 'Agenten – kontrollera källan' (aktuell men sidan "
              "gick inte att läsa), 'Agenten – gammal siffra' och 'Misstänkt fel'. Källan är klickbar och "
              "citatet visar exakt vad som står. Antal deltagare är förifyllt med aktiva där källan skiljer på "
              "aktiva och stödmedlemmar, annars med alla medlemmar."),
        ("p", "Kontaktperson från visar om kontakten kommer från säljarens egen lista (2020–2021) och behöver "
              "bekräftas."),
        ("p", "Sociala medier, Webbshop, Grundår och Om föreningen kommer från föreningens egen webbplats. "
              "Webbshop betyder att föreningen redan säljer kläder eller produkter, en bra öppning i samtalet."),
        ("h2", "Namnen"),
        ("p", "Förening är namnet precis som i originalfilen och ändras aldrig. Matcha händelser mot "
              "originalet med Förening eller Rad i Leadsbas (radnumret i fliken Leadsbas); e-postadressen "
              "är också unik per rad."),
        ("p", "Korrigerat namn är det som står i mejlen: utan parenteser (till exempel '(Lag: P-07)'), "
              "utan anteckningar som '/Konkurs' och med stavningen kontrollerad. Rosa ruta = stavningen är "
              "ändrad, Namnändring säger vad. Skriv över Korrigerat namn om något blev fel."),
        ("p", f"Borttagna: {len(stopped)} föreningar med 'slutat' eller 'upphört' i namnet finns inte med "
              "(listan längst ner). "
              f"{names['dormant']} med 'konkurs', 'vilande', 'inaktiv' eller 'kundförl' i namnet är kvar men "
              "flaggade i Kontrollera."),
        ("p", f"Stavning: {names['spelled']} namn rättade. {names['cleaned']} namn har fått parentes eller "
              "anteckning borttagen."),
        ("h2", "Flikarna"),
        ("p", "Kent – utskick / Niclas – utskick: alla föreningar som är taggade på säljaren, utom lagen. "
              "Importeras i Lead-mejl."),
        ("p", f"{CENTRAL} – intro: föreningar utan säljare som har e-post. Får introduktionen från Johan Lindqvist."),
        ("p", f"{TEAM_SHEET}: {len(teams)} lag, flyttade från säljarflikarna så att de kan behandlas för sig. "
              f"{sum(t['team'].startswith('Ordet') for t in teams)} har ordet lag i namnet (Lag: P-07, Damlaget, "
              f"A-lag), {sum(not t['team'].startswith('Ordet') for t in teams)} har en lagbeteckning utan ordet "
              "(Team P05, P-05, F-06/07, Trupp). Säljare och Varför i Lag-fliken står längst till höger."),
        ("p", "Uppvärmning: hur många mejl per dag, dag för dag, och reglerna för uppföljning."),
        ("p", "Antaganden: belopp per deltagare (6 × 399 kr = 2 394 kr) och 35 % till föreningen."),
        ("h2", "Färgerna i rubrikraden"),
        ("p", "Gul = Kent och Niclas fyller i (tomma gula rutor saknas). Brun = från listan. "
              "Ljusbrun = räknas automatiskt. Grön = fylls i efter utskick (skickat, öppnat, svarat, nästa steg). "
              "Hela raden blir grön när den är klar för det personliga mejlet."),
        ("h2", "Läget"),
        ("p", f"Introduktionen: {intro_total} föreningar med e-post (Kent {k['email']}, Niclas {n['email']}, "
              f"Johan {j['email']}). Varje adress finns bara på en rad."),
        ("p", f"Kent: {k['total']} föreningar, {k['members']} med medlemsantal, {k['ready']} klara för personligt "
              f"mejl ({k['prio1']} med tidigare försäljning = ordning 1)."),
        ("p", f"Niclas: {n['total']} föreningar, varav {n['cat_a']} i kategori A. {n['members']} med medlemsantal, "
              f"{n['ready']} klara för personligt mejl."),
        ("p", f"{len(skipped)} föreningar markerade 'Ta ej med' är inte med. {shared} adresser delas av flera "
              f"föreningar och står bara på en rad. {dead} adresser ligger på domäner som inte tar emot mejl "
              "och är flyttade till Kontrollera."),
        ("h2", "Manuellt jobb för Kent och Niclas"),
        ("p", f"0. Sätt Segment. Först {SEG_DIALOG} på föreningarna ni redan har dialog med, sedan justera "
              "Personligt och Generellt."),
        ("p", "Antal deltagare är förifyllt där det finns en siffra. Bedöm och justera till aktiva medlemmar, "
              "så räknas försäljning och behållning om direkt."),
        ("p", f"1. Niclas: fyll i Antal deltagare för kategori A först ({n['cat_a_need']} rader). "
              f"Kent: börja med ordning 1 och överskott ({k['need_members']} rader med e-post saknar antal)."),
        ("p", f"2. Gå igenom Kontrollera (Kent {k['check']}, Niclas {n['check']}): delade adresser, stora "
              "föreningar och döda adresser. Töm rutan när det är löst."),
        ("p", f"3. Hitta e-post där den saknas (Kent {k['no_email']}, Niclas {n['no_email']}). Nya adresser "
              "före söndag kväll 1 november hinner med från start."),
        ("p", "4. Valfritt: kontaktperson (ger 'Hej Anna!' i stället för 'Hej!') och webbplats "
              "(ger föreningens logga i shopförslaget)."),
        ("p", "Skriv bara i gula kolumner och de gröna uppföljningskolumnerna. Ändra inte Förening eller Rad i Leadsbas, "
              "de behövs för att matcha mot originalfilen."),
        ("h2", "Så används kolumnerna i det personliga mejlet"),
        ("p", "[Föreningsnamn] = Korrigerat namn. [Förnamn] = första namnet i Kontaktperson. "
              "[antal deltagare] = Antal deltagare. [belopp] = Belopp per deltagare. "
              "[total försäljning] = Total försäljning. [Behållning] = Behållning 35 %. "
              "[Se shopförslaget] = länk till föreningens egen shop, med logga från Webbplats."),
    ]
    if skipped:
        lines.append(("h2", "Ta ej med"))
        for seller, name, row in skipped:
            lines.append(("p", f"{name} ({seller}, rad {row} i Leadsbas)"))
    if stopped:
        lines.append(("h2", "Borttagna: slutat eller upphört"))
        for seller, name, row in stopped:
            lines.append(("p", f"{name} ({seller}, rad {row} i Leadsbas)"))

    r = 1
    for kind, t in lines:
        if kind == "h2":
            r += 1
        c = ws.cell(r, 2, t)
        if kind == "h1":
            c.font = Font(bold=True, size=16, color=BROWN)
            ws.row_dimensions[r].height = 26
        elif kind == "h2":
            c.font = Font(bold=True, size=12, color=BROWN)
        else:
            c.alignment = Alignment(wrap_text=True, vertical="top")
            ws.row_dimensions[r].height = 15 * max(1, -(-len(t) // 120))
        r += 1
    ws.sheet_view.showGridLines = False


def main(src: str, dst: str, enrichments: list[str], contacts: list[str], newwave: str | None,
         spelling: str | None):
    wb = load_workbook(src)
    leads, skipped, stopped = collect(wb["Leadsbas"])
    for spec in contacts:
        seller, _, path = spec.partition("=")
        print(f"kontakter {seller}:", apply_contacts(leads, seller, path))
    for path in enrichments:
        print(f"berikning {path}:", apply_enrichment(leads, path))
    if newwave:
        print("newwave:", apply_newwave(leads, newwave))
    shared = mark_shared(leads)
    mark_big(leads)
    dead = mark_dead_domains(leads)
    everyone = [l for s in leads for l in leads[s]]
    names = {
        "spelled": apply_spelling(leads, spelling) if spelling else 0,
        "cleaned": sum(l["clean"] != l["name"] for l in everyone),
        "dormant": sum(bool(DORMANT.search(l["name"])) for l in everyone),
    }
    teams = split_teams(leads)
    st = {}
    for seller in SELLERS:
        title = f"{seller} – intro" if seller == CENTRAL else f"{seller} – utskick"
        write_seller_sheet(wb, title, seller, leads[seller])
        st[seller] = stats(leads[seller])
    write_seller_sheet(wb, TEAM_SHEET, None, teams)
    write_warmup(
        wb,
        {s: st[s]["email"] for s in SELLERS},
        {s: st[s]["ready"] for s in SELLERS},
    )
    write_assumptions(wb)
    write_readme(wb, st, skipped, shared, dead, stopped, teams, names)
    originals = [wb["Leadsbas"], wb["Kontroll"]]
    wb._sheets = [ws for ws in wb.worksheets if ws not in originals] + originals
    wb.active = 0
    for ws in wb.worksheets:
        ws.sheet_view.tabSelected = ws.title == "Läs mig"
    wb.save(dst)
    for seller in SELLERS:
        print(seller, st[seller])
    print("ta ej med:", len(skipped), "delade adresser:", shared, "döda domäner:", dead)
    print("slutat/upphört:", len(stopped), "lag:", len(teams), "namn:", names)


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("src")
    ap.add_argument("dst")
    ap.add_argument("--berikning", action="append", default=[])
    ap.add_argument("--kontakter", action="append", default=[])
    ap.add_argument("--newwave", help="newwave_per_rad.json: Newwaves klubbregister per rad i Leadsbas")
    ap.add_argument("--stavning", help="namn_stavning.json från spellcheck_names.py")
    a = ap.parse_args()
    main(a.src, a.dst, a.berikning, a.kontakter, a.newwave, a.stavning)
